from __future__ import annotations

import json
import os
import re
import secrets
import threading
from functools import wraps
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from urllib.parse import parse_qs, urlparse, urlencode
from io import BytesIO
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, HTTPException, Query, Request, Response
from fastapi.responses import RedirectResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from openpyxl import load_workbook

try:  # Supports both `uvicorn code.main:app` and direct test imports.
    from .rag import retrieve_articles
    from . import salesforce, local_beta
except ImportError:  # pragma: no cover
    from rag import retrieve_articles
    import salesforce
    import local_beta

BASE_DIR = Path(__file__).resolve().parents[1]
# This application is configured from its project-local environment file.  Override
# inherited empty placeholders so a local restart uses the values the operator set.
load_dotenv(BASE_DIR / ".env", override=True)
TICKETS_DIR = BASE_DIR / "support_files"
HELP_CENTER_DIR = BASE_DIR / "data" / "CA"
FRONTEND_DIST = BASE_DIR.parent / "frontend" / "dist"

TEAM_MEMBERS = {
    "alex-rivera": {"id": "alex-rivera", "name": "Alex Rivera", "role": "Support lead", "initials": "AR"},
    "priya-shah": {"id": "priya-shah", "name": "Priya Shah", "role": "Roster specialist", "initials": "PS"},
    "jordan-lee": {"id": "jordan-lee", "name": "Jordan Lee", "role": "Support analyst", "initials": "JL"},
    "sam-chen": {"id": "sam-chen", "name": "Sam Chen", "role": "Escalation manager", "initials": "SC"},
}

ROSTER_CASES = {
    "Login or Missing Entity": ("failed-login-missing-entity", r"\b(login|log in|sign in|sso|missing (user|student|section|enrollment)|entity)\b"),
    "No Classes Assigned": ("not-assigned-to-classes", r"\b(not assigned to any classes|no classes assigned)\b"),
    "Student Reactivation": ("student-reactivation", r"\b(reactiv\w*|student id number is already in use)\b"),
    "User Deactivation": ("user-deactivation", r"\b(deactiv\w*|zz_?closed|to be deactivated)\b"),
    "Filter Change": ("filter-change", r"\b(filter|add|remove|subject|course|school|grade)\b"),
    "Student ID Conversion": ("student-id-conversion", r"\b(student id|id conversion)\b"),
    "Push or Sync": ("push-sync", r"\b(push|sync|model)\b"),
    "Threshold Override": ("threshold-overrides", r"\b(threshold|nightly job)\b"),
    "Password Reset": ("password-reset", r"\b(password reset|reset password|forgot password|locked out)\b"),
    "Credential Update": ("credential-updates", r"\b(credential|username|auto_)\b"),
    "Admin Account Creation": ("admin-account-creation", r"\b(create|new)\b.*\badmin\b|\badmin\b.*\b(create|new)\b"),
    "Syncing Admin Behavior": ("syncing-admin-behavior", r"\b(syncing admins?|admin sync)\b"),
    "i-Ready Inform Recreate": ("inform-recreate", r"\b(inform|historical data|wrong record)\b"),
    "Demographics": ("demographics", r"\b(demographic|race|gender|ell|iep)\b"),
}
DEFAULT_CATEGORIES = [*ROSTER_CASES, "Needs Routing"]

HUMAN_PATTERNS = {
    "security or privacy concern": r"\b(hack|hacked|breach|fraud|unauthori[sz]ed|security|privacy|gdpr)\b",
    "payment dispute": r"\b(chargeback|refund|charged twice|payment dispute|invoice dispute)\b",
    "data-loss risk": r"\b(data loss|deleted .*data|lost .*data|restore .*data)\b",
    "urgent request": r"\b(urgent|asap|immediately|production down|outage|critical)\b",
    "sensitive conduct report": r"\b(harass\w*|abus\w*|threat\w*|discriminat\w*)\b",
}


class TicketCreate(BaseModel):
    subject: str = Field(min_length=3, max_length=160)
    customer: str = Field(min_length=2, max_length=80)
    message: str = Field(min_length=5, max_length=5000)


class TicketUpdate(BaseModel):
    status: Literal["open", "in_progress", "resolved", "escalated"] | None = None
    owner: str | None = Field(default=None, max_length=80)
    resolution_note: str | None = Field(default=None, max_length=3000)


class SuggestionRequest(BaseModel):
    agent_notes: str | None = Field(default=None, max_length=2000)


class SalesforceApply(BaseModel):
    action: Literal["mark_human", "mark_resolved"]


class SalesforceSyncRequest(BaseModel):
    queue_id: str = Field(min_length=15, max_length=18)


class TeamSignIn(BaseModel):
    member_id: str


def now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def request_sessions(request: Request, *, local_only: bool = False) -> list:
    sessions = []
    if not local_only and sf_session_id(request):
        try:
            sessions.append(salesforce._session(sf_session_id(request)))
        except HTTPException as error:
            if error.status_code != 401:
                raise
    beta = local_beta.lookup(request)
    if beta:
        sessions.append(beta)
    if not sessions:
        raise HTTPException(401, "Start an Excel beta session or connect Salesforce before continuing.")
    return sessions


def require_member(request: Request, ticket: dict | None = None) -> dict:
    sessions = request_sessions(request)
    if ticket and ticket.get("salesforce_instance_url") == local_beta.SCOPE:
        sessions = [session for session in sessions if session.instance_url == local_beta.SCOPE]
    session = sessions[0]
    return {"id": session.user_id, "name": "Local beta agent" if session.instance_url == local_beta.SCOPE else salesforce.display_name(session), "role": "Support agent"}


def visible_tickets(request: Request, *, local_only: bool = False) -> list[dict]:
    sessions = request_sessions(request, local_only=local_only)
    tickets = all_tickets()
    visible = set()
    for session in sessions:
        if session.instance_url != local_beta.SCOPE:
            candidates = [ticket for ticket in tickets if ticket.get("source") == "salesforce" and ticket.get("salesforce_instance_url") == session.instance_url]
            allowed = salesforce.accessible_case_ids(session, [ticket["salesforce_case_id"] for ticket in candidates])
            visible.update(ticket["id"] for ticket in candidates if ticket["salesforce_case_id"] in allowed)
        visible.update(ticket["id"] for ticket in tickets if ticket.get("source") != "salesforce" and ticket.get("created_by") == session.user_id and ticket.get("salesforce_instance_url") == session.instance_url)
    return [ticket for ticket in tickets if ticket["id"] in visible]


def authorized_ticket(ticket_id: str, request: Request) -> dict:
    sessions = request_sessions(request)
    ticket = load_ticket(ticket_id)
    allowed = False
    for session in sessions:
        if ticket.get("salesforce_instance_url") != session.instance_url:
            continue
        if ticket.get("source") == "salesforce" and session.instance_url != local_beta.SCOPE:
            allowed = ticket.get("salesforce_case_id") in salesforce.accessible_case_ids(session, [ticket.get("salesforce_case_id", "")])
        elif ticket.get("source") != "salesforce":
            allowed = ticket.get("created_by") == session.user_id
        if allowed:
            break
    if not allowed:
        raise HTTPException(404, "Ticket not found")
    return ticket


def record_activity(ticket: dict, member: dict, action: str, details: dict | None = None) -> None:
    ticket.setdefault("activity", []).append({
        "at": now(), "member_id": member["id"], "member_name": member["name"],
        "action": action, "details": details or {},
    })


def case_reason(ticket: dict) -> str | None:
    """Return the source Case Reason when the case came from an export."""
    return ticket.get("case_reason") or ticket.get("imported_case_fields", {}).get("Case Reason")


def load_articles() -> list[dict]:
    articles = []
    paths = sorted({*HELP_CENTER_DIR.glob("*.md"), *HELP_CENTER_DIR.glob("*.md.txt")})
    for path in paths:
        raw = path.read_text(encoding="utf-8")
        first_line = next((line[2:].strip() for line in raw.splitlines() if line.startswith("# ")), path.name.removesuffix(".md.txt").removesuffix(".md"))
        articles.append({"slug": path.name.removesuffix(".md.txt").removesuffix(".md"), "title": first_line, "content": raw})
    return articles


def classify(text: str) -> str:
    lowered = text.lower()
    candidates: list[tuple[int, str]] = []
    for category, (_, pattern) in ROSTER_CASES.items():
        match = re.search(pattern, lowered)
        if match:
            # Prefer the most specific matching phrase over broad terms such as
            # "sync", "school", or "grade" that appear in multiple case types.
            candidates.append((len(match.group(0)), category))
    return max(candidates, default=(0, "Needs Routing"))[1]


def find_article(category: str, text: str, articles: list[dict]) -> dict | None:
    case = ROSTER_CASES.get(category)
    return next((article for article in articles if case and article["slug"] == case[0]), None)


def category_for_slug(slug: str) -> str:
    return next((category for category, (case_slug, _) in ROSTER_CASES.items() if case_slug == slug), "Needs Routing")


def select_playbook(subject: str, message: str) -> tuple[str, dict | None, list[dict], float, str]:
    """Plan retrieval in two stages: precise case signal, then KB retrieval."""
    query = f"{subject} {message}"
    retrieved = retrieve_articles(HELP_CENTER_DIR, query)
    classified_category = classify(query)
    known_article = find_article(classified_category, query, load_articles())
    if known_article:
        retrieved_match = next((item for item in retrieved if item["slug"] == known_article["slug"]), None)
        article = retrieved_match or known_article
        return classified_category, article, retrieved, 0.92 if retrieved_match and retrieved_match.get("matched_aliases") else 0.82, "Matched the Roster Support case router."
    if retrieved:
        top = retrieved[0]
        # Without a router match, only select a playbook when a declared alias
        # matches or the title/case type has clear evidence. Otherwise ask to route.
        strong_match = bool(top.get("matched_aliases")) or top.get("raw_score", 0) >= 8
        if strong_match:
            return category_for_slug(top["slug"]), top, retrieved, 0.74, "Matched the knowledge-base case metadata."
    return "Needs Routing", None, retrieved, 0.35, "No sufficiently specific case-playbook match was found."


def playbook_metadata(content: str) -> dict[str, str]:
    """Read the small YAML header used by the supplied case playbooks."""
    if not content.startswith("---"):
        return {}
    header = content.split("---", 2)[1]
    return {
        key.strip(): value.strip().strip('"')
        for line in header.splitlines() if ":" in line
        for key, value in [line.split(":", 1)]
    }


def playbook_steps(content: str, limit: int = 3) -> list[str]:
    """Extract first actions only from named procedure sections, not tables."""
    steps: list[str] = []
    in_action_section = False
    for line in content.splitlines():
        cleaned = line.strip()
        if cleaned.startswith("## "):
            heading = cleaned[3:].lower()
            in_action_section = any(key in heading for key in ("before", "step", "intake", "eligibility", "pick the right", "first", "path"))
            continue
        if in_action_section and re.match(r"^(?:[-*]|\d+\.)\s+", cleaned):
            step = re.sub(r"^(?:[-*]|\d+\.)\s+", "", cleaned)
            step = re.sub(r"[*`]+", "", step).strip()
            if len(step) >= 18 and step not in steps and not step.lower().startswith("path "):
                steps.append(step)
        if len(steps) == limit:
            break
    return steps


def build_suggestion(article: dict | None) -> tuple[str | None, str, str, bool]:
    """Return suggested approach, priority, reason, and escalation status."""
    if not article:
        return (
            "Clarify whether this is a roster, sync, provisioning, or access case. If it is outside Roster Support, route it to the appropriate team.",
            "medium",
            "No matching Roster Support case playbook was found, so the case needs routing or clarification.",
            False,
        )

    metadata = playbook_metadata(article["content"])
    risk = metadata.get("risk", "medium").lower()
    priority = {"low": "low", "medium": "medium", "high": "high"}.get(risk, "medium")
    destructive = metadata.get("destructive_steps")
    sensitive = metadata.get("data_sensitivity") == "high"
    owner = metadata.get("owner")
    unavailable = "[Procedure not available]" in article["content"]
    needs_escalation = bool(owner or unavailable)
    if destructive or sensitive:
        priority = "high" if priority != "low" else "medium"
    if needs_escalation:
        priority = "high"

    approach = [
        "Review Provisioning Notes; confirm the account, sync source, requester authority, SSO mode, and affected-user count.",
        *playbook_steps(article["content"]),
    ]
    if destructive:
        approach.append("Prepare the change only after written authorization and explicit confirmation of scope; an employee must approve or execute it.")
    if sensitive:
        approach.append("Treat demographic and student/staff information as PII; use only approved internal systems and the minimum necessary data.")
    if unavailable:
        approach.append("Stop: this playbook says the required procedure is unavailable. Do not reconstruct or guess the missing process.")
    if owner:
        approach.append(f"Route or coordinate with the documented owner: {owner}.")

    reason_parts = [f"{article['title']} is marked {risk} risk."]
    if destructive:
        reason_parts.append(f"It includes controlled action(s): {destructive}.")
    if sensitive:
        reason_parts.append("It involves sensitive data handling.")
    if unavailable:
        reason_parts.append("The required procedure is unavailable, so this requires escalation.")
    if owner:
        reason_parts.append(f"The playbook assigns ownership to {owner}.")
    return "\n".join(f"{index}. {step}" for index, step in enumerate(approach, start=1)), priority, " ".join(reason_parts), needs_escalation


def triage(subject: str, message: str, *, category: str | None = None, article: dict | None = None, confidence: float | None = None, retrieval_reason: str | None = None) -> dict:
    text = f"{subject} {message}"
    category = category or classify(text)
    article = article or find_article(category, text, load_articles())
    matches = [reason for reason, pattern in HUMAN_PATTERNS.items() if re.search(pattern, text, re.I)]
    suggestion, priority, playbook_reason, needs_escalation = build_suggestion(article)
    if article is None or (confidence is not None and confidence < 0.7):
        needs_escalation = True
        playbook_reason += " Low-confidence requests require human escalation."

    if matches:
        priority = "critical" if any(reason in matches for reason in ["security or privacy concern", "data-loss risk", "urgent request"]) else "high"
        return {
            "category": category,
            "priority": priority,
            "status": "escalated",
            "requires_human": True,
            "confidence": 0.96,
            "rationale": f"Escalated because it contains {', '.join(matches)}. {retrieval_reason or playbook_reason} {playbook_reason}",
            "recommended_article": article["title"] if article else None,
            "suggested_approach": suggestion,
            "auto_resolution": None,
        }

    return {
        "category": category,
        "priority": priority,
        "status": "escalated" if needs_escalation else "open",
        "requires_human": True,
        "confidence": confidence if confidence is not None else (0.76 if article else 0.42),
        "rationale": f"{retrieval_reason or ''} {playbook_reason}".strip(),
        "recommended_article": article["title"] if article else None,
        "suggested_approach": suggestion,
        "auto_resolution": None,
    }


def rag_triage(subject: str, message: str, agent_notes: str | None = None) -> dict:
    """Retrieve help-center context before applying the safe local triage policy.

    An OpenAI key may be configured later for drafting, but the deterministic
    safety gate remains the authority for automatic closure and escalation.
    """
    category, article, matches, confidence, retrieval_reason = select_playbook(subject, message)
    result = triage(subject, message, category=category, article=article, confidence=confidence, retrieval_reason=retrieval_reason)
    result["rag_articles"] = [{
        "title": item["title"], "slug": item["slug"], "score": item["score"],
        "matched_aliases": item.get("matched_aliases", []),
    } for item in matches]
    result["ai_mode"] = "rule_based_rag"
    # An agent explicitly requested this suggestion, so optional model assist
    # may refine the approach and provide an advisory severity assessment.
    # The deterministic CA/safety priority remains the queue authority.
    if os.getenv("OPENAI_API_KEY") and os.getenv("ENABLE_OPENAI_TRIAGE", "false").lower() == "true":
        result = apply_model_assist(result, subject, message, matches, agent_notes)
    return result


def pending_approach() -> dict:
    """Safe initial state until an agent explicitly requests a suggestion."""
    return {
        "category": "Needs Routing",
        "priority": "medium",
        "status": "open",
        "requires_human": True,
        "confidence": 0.0,
        "rationale": "No approach has been suggested yet. Open this case and select 'Suggest me the approach'.",
        "recommended_article": None,
        "suggested_approach": None,
        "auto_resolution": None,
        "rag_articles": [],
        "ai_mode": "not_requested",
        "ai_advisory_priority": None,
        "ai_priority_rationale": None,
    }


def apply_model_assist(result: dict, subject: str, message: str, matches: list[dict], agent_notes: str | None) -> dict:
    """Refine a suggestion and return advisory priority; deterministic gates win."""
    try:
        from openai import OpenAI

        schema = {
            "type": "object",
            "properties": {
                "refined_suggestion": {"type": "string", "maxLength": 2200},
                "advisory_priority": {"type": "string", "enum": ["low", "medium", "high", "critical"]},
                "priority_rationale": {"type": "string", "maxLength": 500},
            },
            "required": ["refined_suggestion", "advisory_priority", "priority_rationale"],
            "additionalProperties": False,
        }
        context = "\n\n".join(f"[{article['title']}]\n{article['content']}" for article in matches)
        response = OpenAI().responses.create(
            model=os.getenv("OPENAI_MODEL", "gpt-5-mini"),
            store=False,
            instructions=("Refine the supplied local knowledge-base approach for an internal support agent and assess case urgency. "
                          "Use only the supplied playbooks. Keep it concise and actionable. Preserve every stop, escalation, authorization, PII, and human-review requirement. "
                          "Never invent links, internal steps, facts, or outcomes; if the playbook is incomplete, explicitly say to stop and escalate. "
                          "The priority is advisory only: assess the customer impact and urgency in the case, but do not claim it overrides the supplied deterministic safety priority."),
            input=(f"Case subject: {subject}\nCase message: {message}\n\n"
                   f"Agent refinement notes: {agent_notes or 'None provided.'}\n\n"
                   f"Deterministic safety priority (authoritative for the queue): {result['priority']}\n\n"
                   f"Current deterministic approach:\n{result['suggested_approach']}\n\n"
                   f"Help-center context:\n{context}"),
            text={"format": {"type": "json_schema", "name": "ticket_triage", "strict": True, "schema": schema}},
        )
        model_result = json.loads(response.output_text)
        refined = model_result["refined_suggestion"].strip()
        if refined:
            result["suggested_approach"] = refined
            result["rationale"] = f"{result['rationale']} AI refinement was applied using the retrieved playbooks."
        result["ai_advisory_priority"] = model_result["advisory_priority"]
        result["ai_priority_rationale"] = model_result["priority_rationale"].strip()
        result["ai_mode"] = "openai_rag_refined"
    except Exception as exc:
        result["ai_mode"] = "rule_based_rag_fallback"
        result["ai_advisory_priority"] = None
        result["ai_priority_rationale"] = None
        result["rationale"] = f"{result['rationale']} Model assistance was unavailable; the safe local policy was used."
    return result


# The supported launcher runs one worker. Serialize all ticket mutations and
# atomically replace files so readers never observe partially written JSON.
TICKET_LOCK = threading.RLock()


def serialized_ticket_write(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        with TICKET_LOCK:
            return function(*args, **kwargs)
    return wrapped


def ticket_path(ticket_id: str) -> Path:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", ticket_id):
        raise HTTPException(400, "Invalid ticket ID")
    return TICKETS_DIR / f"{ticket_id}.json"


def load_ticket(ticket_id: str) -> dict:
    path = ticket_path(ticket_id)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Ticket not found")
    return json.loads(path.read_text(encoding="utf-8"))


def save_ticket(ticket: dict) -> dict:
    TICKETS_DIR.mkdir(parents=True, exist_ok=True)
    path = ticket_path(ticket["id"])
    temporary = path.with_suffix(".json.tmp")
    with TICKET_LOCK:
        temporary.write_text(json.dumps(ticket, indent=2), encoding="utf-8")
        temporary.replace(path)
    return ticket


def import_excel_tickets(contents: bytes, *, created_by: str | None = None, instance_url: str | None = None) -> tuple[list[dict], list[str]]:
    """Create local, untriaged tickets from a beta Excel import."""
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Upload a valid .xlsx workbook with a header row.") from exc

    if not headers:
        raise HTTPException(status_code=400, detail="The workbook is empty.")
    normalized = {str(value).strip().lower(): index for index, value in enumerate(headers) if value is not None}
    fields = {
        "subject": next((normalized[name] for name in ("subject", "title") if name in normalized), None),
        "customer": next((normalized[name] for name in ("customer", "customer name", "contact", "name", "account name") if name in normalized), None),
        "message": next((normalized[name] for name in ("message", "description", "case description", "details") if name in normalized), None),
    }
    missing = [name for name in ("subject", "customer") if fields[name] is None]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing required column(s): {', '.join(missing)}. Use Subject and Customer, or Subject and Account Name for a Case export.")

    imported: list[dict] = []
    skipped: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        if len(imported) >= 200:
            skipped.append("Only the first 200 valid rows were imported.")
            break
        row_values = {
            str(header).strip(): str(row[index]).strip() if index < len(row) and row[index] is not None else ""
            for index, header in enumerate(headers) if header is not None
        }
        values = {name: str(row[index]).strip() if index is not None and index < len(row) and row[index] is not None else "" for name, index in fields.items()}
        if not values["message"]:
            case_details = [
                f"{label}: {row_values[label]}"
                for label in ("Case Number", "Case Reason", "Case Sub Reason", "Product", "State", "Segment", "Priority", "Case Owner")
                if row_values.get(label)
            ]
            values["message"] = "Imported case details: " + "; ".join(case_details)
        if not any(values.values()):
            continue
        if not (3 <= len(values["subject"]) <= 160 and 2 <= len(values["customer"]) <= 80 and 5 <= len(values["message"]) <= 5000):
            skipped.append(f"Row {row_number} was skipped: Subject (3-160), Customer (2-80), and Message (5-5000) are required.")
            continue
        ticket_id = f"IMP-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')}-{row_number}"
        ticket = {
            "id": ticket_id,
            "source": "excel_beta_import",
            "created_by": created_by,
            "salesforce_instance_url": instance_url,
            "imported_case_fields": {key: value for key, value in row_values.items() if value},
            "case_reason": row_values.get("Case Reason") or None,
            **values,
            "owner": None,
            "resolution_note": None,
            "created_at": now(),
            "updated_at": now(),
            **pending_approach(),
        }
        imported.append(save_ticket(ticket))
    return imported, skipped


def all_tickets() -> list[dict]:
    TICKETS_DIR.mkdir(parents=True, exist_ok=True)
    tickets = [json.loads(path.read_text(encoding="utf-8")) for path in TICKETS_DIR.glob("*.json")]
    return sorted(tickets, key=lambda ticket: (ticket["priority"] != "critical", ticket["priority"] != "high", ticket["created_at"]), reverse=False)


FRONTEND_URL = os.getenv("FRONTEND_URL", "http://127.0.0.1:5173").rstrip("/")

app = FastAPI(title="Support Triage API", version="1.0.0")


@app.middleware("http")
async def prevent_private_response_caching(request: Request, call_next):
    response = await call_next(request)
    if request.url.path.startswith(("/tickets", "/cases", "/team", "/me/", "/knowledge-base", "/auth/")):
        response.headers["Cache-Control"] = "no-store"
    return response


app.add_middleware(
    CORSMiddleware,
    # Vite can be opened through either host locally. The frontend defaults to
    # 127.0.0.1 for its API, so allow its matching browser origin as well.
    allow_origins=["https://localhost:5173", FRONTEND_URL],
    allow_methods=["*"],
    allow_headers=["*"],
    allow_credentials=True,
)


@app.get("/health")
def health() -> dict:
    return {"status": "ok"}


@app.get("/tickets")
def list_tickets(
    request: Request,
    local_only: bool = False,
    status: str | None = None,
    category: str | None = None,
    priority: str | None = None,
    case_reason_filter: str | None = Query(default=None, alias="case_reason", max_length=120),
    search: str | None = Query(default=None, max_length=160),
) -> list[dict]:
    tickets = visible_tickets(request, local_only=local_only)
    # Backfill the response field for imports created before `case_reason` was
    # promoted from the retained source fields, without changing old files.
    for ticket in tickets:
        ticket["case_reason"] = case_reason(ticket)
    if status:
        tickets = [ticket for ticket in tickets if ticket["status"] == status]
    if category:
        tickets = [ticket for ticket in tickets if ticket["category"] == category]
    if priority:
        tickets = [ticket for ticket in tickets if ticket["priority"] == priority]
    if case_reason_filter:
        tickets = [ticket for ticket in tickets if case_reason(ticket) == case_reason_filter]
    if search:
        needle = search.lower()
        tickets = [ticket for ticket in tickets if needle in f"{ticket['subject']} {ticket['customer']} {ticket['message']}".lower()]
    return tickets


@app.get("/tickets/{ticket_id}")
def get_ticket(ticket_id: str, request: Request) -> dict:
    return authorized_ticket(ticket_id, request)


@app.post("/tickets/import/excel")
async def import_excel(request: Request) -> dict:
    session = request_sessions(request)[0]
    contents = await request.body()
    if not contents:
        raise HTTPException(status_code=400, detail="Choose an .xlsx file to import.")
    imported, skipped = import_excel_tickets(contents, created_by=session.user_id, instance_url=session.instance_url)
    return {"imported": len(imported), "tickets": imported, "skipped": skipped}


@app.post("/tickets", status_code=201)
def create_ticket(payload: TicketCreate, request: Request) -> dict:
    session = request_sessions(request)[0]
    ticket_id = f"TKT-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')}"
    ticket = {
        "id": ticket_id,
        "created_by": session.user_id,
        "salesforce_instance_url": session.instance_url,
        "subject": payload.subject,
        "customer": payload.customer,
        "message": payload.message,
        "owner": None,
        "resolution_note": None,
        "created_at": now(),
        "updated_at": now(),
        **pending_approach(),
    }
    return save_ticket(ticket)


@app.post("/tickets/{ticket_id}/suggest-approach")
@serialized_ticket_write
def suggest_ticket_approach(ticket_id: str, payload: SuggestionRequest, request: Request) -> dict:
    """Run retrieval and triage only after an agent requests an approach."""
    ticket = authorized_ticket(ticket_id, request)
    member = require_member(request, ticket)
    result = rag_triage(ticket["subject"], ticket["message"], payload.agent_notes)
    workflow_status = ticket["status"]
    ticket.update(result)
    if result["status"] != "escalated":
        ticket["status"] = workflow_status
    record_activity(ticket, member, "requested_suggestion", {"used_ai_refinement": bool(payload.agent_notes)})
    ticket["updated_at"] = now()
    return save_ticket(ticket)


@app.patch("/tickets/{ticket_id}")
@serialized_ticket_write
def update_ticket(ticket_id: str, payload: TicketUpdate, request: Request) -> dict:
    ticket = authorized_ticket(ticket_id, request)
    member = require_member(request, ticket)
    changes = payload.model_dump(exclude_none=True)
    if changes.get("status") == "resolved" and ticket.get("status") != "resolved":
        ticket.setdefault("resolution_events", []).append({
            "at": now(), "member_id": member["id"],
            "instance_url": ticket.get("salesforce_instance_url"),
        })
    ticket.update(changes)
    record_activity(ticket, member, "updated_case", {"fields": sorted(changes)})
    ticket["updated_at"] = now()
    return save_ticket(ticket)


@app.get("/categories")
def categories() -> list[str]:
    return DEFAULT_CATEGORIES


@app.get("/knowledge-base")
def knowledge_base(request: Request) -> list[dict]:
    """Expose the supplied CA articles for agents to browse in the UI."""
    require_member(request)
    return [{"slug": article["slug"], "title": article["title"], "content": article["content"]} for article in load_articles()]


@app.get("/team/members")
def team_members(request: Request) -> list[dict]:
    return [require_member(request)]


@app.post("/team/sign-in")
def team_sign_in(payload: TeamSignIn, request: Request) -> dict:
    member = require_member(request)
    if payload.member_id != member["id"]:
        raise HTTPException(403, "Use your authenticated identity.")
    return member


@app.get("/me/progress")
def my_progress(request: Request, days: int = 30, timezone_name: str = Query("UTC", alias="timezone")) -> dict:
    session = salesforce._session(sf_session_id(request))
    if days not in (7, 30, 90):
        raise HTTPException(422, "Choose 7, 30, or 90 days.")
    try:
        zone = ZoneInfo(timezone_name)
    except (ZoneInfoNotFoundError, ValueError):
        raise HTTPException(422, "Choose a valid IANA timezone.")
    today = datetime.now(timezone.utc).astimezone(zone).date()
    start = today - timedelta(days=days - 1)
    counts = {(start + timedelta(days=i)).isoformat(): 0 for i in range(days)}
    total = 0
    for ticket in visible_tickets(request):
        if ticket.get("salesforce_instance_url") != session.instance_url:
            continue
        for event in ticket.get("resolution_events", []):
            if event.get("member_id") != session.user_id or event.get("instance_url") != session.instance_url:
                continue
            try:
                timestamp = datetime.fromisoformat(event["at"])
                if timestamp.tzinfo is None:
                    continue
                day = timestamp.astimezone(zone).date().isoformat()
            except (KeyError, TypeError, ValueError):
                continue
            total += 1
            if day in counts:
                counts[day] += 1
    return {"total_resolutions": total, "period_resolutions": sum(counts.values()),
            "timezone": timezone_name, "days": days,
            "daily": [{"date": day, "count": count} for day, count in counts.items()]}


@app.get("/team/performance")
def team_performance(request: Request, local_only: bool = False) -> list[dict]:
    sessions = request_sessions(request, local_only=local_only)
    performance = {session.user_id: {"id": session.user_id, "name": "Local beta agent" if session.instance_url == local_beta.SCOPE else salesforce.display_name(session), "case_updates": 0, "suggestions_requested": 0, "cases_touched": set(), "last_activity": None} for session in sessions}
    for ticket in visible_tickets(request, local_only=local_only):
        for activity in ticket.get("activity", []):
            member = performance.get(activity.get("member_id"))
            if not member:
                continue
            member["cases_touched"].add(ticket["id"])
            if activity.get("action") == "requested_suggestion":
                member["suggestions_requested"] += 1
            if activity.get("action") == "updated_case":
                member["case_updates"] += 1
            if not member["last_activity"] or activity.get("at", "") > member["last_activity"]:
                member["last_activity"] = activity.get("at")
    return [{**member, "cases_touched": len(member["cases_touched"])} for member in performance.values()]


def sf_session_id(request: Request) -> str | None:
    return request.cookies.get(salesforce.COOKIE_NAME)


@serialized_ticket_write
def mirror_salesforce_case(case: dict, instance_url: str) -> dict:
    """Persist only the authorized Case data required by this dashboard."""
    case_id = case["Id"]
    mirror_id = f"SF-{salesforce._hash(instance_url)[:16]}-{case_id}"
    path = ticket_path(mirror_id)
    existing = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    ticket = {
        **pending_approach(), **existing,
        "id": mirror_id, "source": "salesforce", "salesforce_case_id": case_id,
        "salesforce_instance_url": instance_url,
        "salesforce_case_number": case.get("CaseNumber"), "subject": case.get("Subject") or f"Salesforce case {case.get('CaseNumber', case_id)}",
        "customer": "Salesforce case", "message": case.get("Description") or "No case description was supplied.",
        "salesforce_status": case.get("Status"), "salesforce_priority": case.get("Priority"),
        "salesforce_case_reason": case.get("CaseReason"), "salesforce_type": case.get("Type"),
        "salesforce_owner_id": case.get("OwnerId"), "salesforce_created_at": case.get("CreatedDate"),
        "salesforce_last_modified_at": case.get("LastModifiedDate"), "case_reason": case.get("CaseReason") or existing.get("case_reason"),
        "owner": existing.get("owner"), "resolution_note": existing.get("resolution_note"),
        "created_at": existing.get("created_at", case.get("CreatedDate") or now()), "updated_at": now(),
    }
    return save_ticket(ticket)


@app.post("/auth/local-beta")
def start_local_beta(request: Request, response: Response) -> dict:
    token = local_beta.create(request)
    response.set_cookie(local_beta.COOKIE_NAME, token, httponly=True, secure=request.url.scheme == "https", samesite="strict", max_age=salesforce.SESSION_SECONDS)
    return {"ready": True}


@app.get("/auth/salesforce/diagnostics")
def salesforce_diagnostics(request: Request) -> dict:
    local_beta.require_loopback(request)
    return {
        "configured": salesforce.configured(),
        "login_url": salesforce.login_url(),
        "callback_url": os.getenv("SALESFORCE_REDIRECT_URI", ""),
        "scopes": salesforce.oauth_scopes().split(),
        "pkce": "S256", "writeback_enabled": salesforce.writeback_enabled(),
        "org_approval_verified": False,
        "approval_help": salesforce.oauth_error_message("OAUTH_APPROVAL_ERROR_GENERIC"),
    }


@app.get("/auth/salesforce")
def salesforce_authorize(return_to: str | None = None) -> RedirectResponse:
    if return_to and return_to not in {FRONTEND_URL, "https://localhost:5173"}:
        raise HTTPException(400, "Unsupported frontend return address.")
    url = salesforce.authorization_url()
    response = RedirectResponse(url)
    response.set_cookie("rqc_oauth_return", return_to or FRONTEND_URL, httponly=True, secure=FRONTEND_URL.startswith("https://"), samesite="lax", max_age=salesforce.OAUTH_SECONDS)
    response.set_cookie(salesforce.STATE_COOKIE, parse_qs(urlparse(url).query)["state"][0], httponly=True, secure=FRONTEND_URL.startswith("https://"), samesite="lax", max_age=salesforce.OAUTH_SECONDS)
    return response


@app.get("/auth/salesforce/callback")
def salesforce_callback(request: Request, state: str = "", code: str | None = None, error: str | None = None, error_description: str = "") -> RedirectResponse:
    if not secrets.compare_digest(request.cookies.get(salesforce.STATE_COOKIE, ""), state) or not state:
        raise HTTPException(400, "Authorization must finish in the browser that started it.")
    destination = request.cookies.get("rqc_oauth_return", FRONTEND_URL)
    if destination not in {FRONTEND_URL, "https://localhost:5173"}:
        destination = FRONTEND_URL
    if error:
        salesforce.discard_authorization(state)
        message = salesforce.oauth_error_message(error, error_description)
        response = RedirectResponse(f"{destination}/?{urlencode({'salesforce_error': message})}")
        response.delete_cookie("rqc_oauth_return")
        response.delete_cookie(salesforce.STATE_COOKIE)
        return response
    if not code:
        raise HTTPException(400, "Salesforce did not return an authorization code. Connect again.")
    session_id = salesforce.complete_authorization(code, state)
    response = RedirectResponse(f"{destination}/?salesforce=connected")
    response.set_cookie(salesforce.COOKIE_NAME, session_id, httponly=True, secure=os.getenv("COOKIE_SECURE", "false").lower() == "true", samesite="lax", max_age=60 * 60 * 8)
    response.delete_cookie("rqc_oauth_return")
    response.delete_cookie(salesforce.STATE_COOKIE)
    return response


@app.get("/auth/salesforce/status")
def salesforce_status(request: Request) -> dict:
    return salesforce.status(sf_session_id(request))


@app.post("/auth/salesforce/sync")
def salesforce_sync(payload: SalesforceSyncRequest, request: Request) -> dict:
    session = salesforce._session(sf_session_id(request))
    records = salesforce.sync_cases(sf_session_id(request), payload.queue_id)
    tickets = [mirror_salesforce_case(case, session.instance_url) for case in records]
    return {"synced": len(tickets), "cases": tickets, "last_synced_at": now()}


@app.post("/auth/salesforce/roster-support/sync")
def sync_roster_support(request: Request) -> dict:
    session = salesforce._session(sf_session_id(request))
    view, records = salesforce.roster_support_cases(session)
    tickets = [mirror_salesforce_case(case, session.instance_url) for case in records]
    synced_at = now()
    with salesforce._connect() as connection:
        connection.execute("UPDATE salesforce_sessions SET last_synced_at=?,last_sync_count=? WHERE session_hash=?", (synced_at, len(tickets), salesforce._hash(session.session_id)))
    return {"state": "found" if view else "missing", "view": view,
            "cases": tickets, "count": len(tickets), "last_synced_at": synced_at}


@app.post("/auth/salesforce/disconnect")
def salesforce_disconnect(request: Request) -> Response:
    salesforce.disconnect(sf_session_id(request))
    response = Response(status_code=204)
    response.delete_cookie(salesforce.COOKIE_NAME)
    return response


@app.get("/cases")
def list_salesforce_cases(
    request: Request,
    status: str | None = None, priority: str | None = None, owner: str | None = None,
    case_reason: str | None = None, search: str | None = Query(default=None, max_length=160),
) -> list[dict]:
    cases = [ticket for ticket in visible_tickets(request) if ticket.get("source") == "salesforce"]
    if status: cases = [case for case in cases if case.get("salesforce_status") == status]
    if priority: cases = [case for case in cases if case.get("salesforce_priority") == priority]
    if owner: cases = [case for case in cases if case.get("salesforce_owner_id") == owner]
    if case_reason: cases = [case for case in cases if case.get("salesforce_case_reason") == case_reason]
    if search:
        needle = search.lower()
        cases = [case for case in cases if needle in f"{case.get('salesforce_case_number', '')} {case['subject']} {case['message']}".lower()]
    return cases


@app.post("/cases/{case_id}/assign-me")
@serialized_ticket_write
def assign_case_to_me(case_id: str, request: Request) -> dict:
    session = salesforce._session(sf_session_id(request))
    ticket = authorized_ticket(f"SF-{salesforce._hash(session.instance_url)[:16]}-{case_id}", request)
    if ticket.get("source") != "salesforce":
        raise HTTPException(status_code=404, detail="Salesforce Case not found.")
    ticket["salesforce_owner_id"] = salesforce.assign_to_me(sf_session_id(request), case_id)
    ticket["updated_at"] = now()
    return save_ticket(ticket)


# Serve the production frontend from the same HTTPS origin as the API. This
# keeps the OAuth callback and HTTP-only session cookie first-party in local use.
if FRONTEND_DIST.is_dir():
    app.mount("/", StaticFiles(directory=FRONTEND_DIST, html=True), name="frontend")
